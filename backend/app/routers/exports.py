"""
ConstructMind AI - Export API Router
Created by Moawia Husnain | Civil Engineer | UET Taxila | +923266915744

Provides endpoints for generating professional Excel reports and XER files
from project schedule data, BOQ, resources, and activities.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, date, timedelta
from typing import List, Dict, Any, Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.project import (
    Project, Activity, BOQItem, Resource, ResourceAssignment,
    WBS, Relationship
)
from app.config import settings

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

router = APIRouter()
logger = logging.getLogger("constructmind.routers.exports")

# ══════════════════════════════════════════════════════════════════
# STYLE CONSTANTS
# ══════════════════════════════════════════════════════════════════

NAVY = "0D1B2A"
DARK_BLUE = "1B2838"
MID_BLUE = "1B4965"
TEAL = "0D7C66"
GOLD = "D4A843"
LIGHT_GOLD = "F5E6CC"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"
VERY_LIGHT_GRAY = "F8F9FA"
MED_GRAY = "E0E0E0"
DARK_GRAY = "333333"
RED = "DC3545"
GREEN = "28A745"
LIGHT_GREEN = "E8F5E9"
LIGHT_RED = "FFEBEE"
LIGHT_YELLOW = "FFF8E1"
LIGHT_BLUE = "E3F2FD"

thin_border = Border(
    left=Side(style='thin', color=MED_GRAY),
    right=Side(style='thin', color=MED_GRAY),
    top=Side(style='thin', color=MED_GRAY),
    bottom=Side(style='thin', color=MED_GRAY),
)

thick_bottom = Border(
    left=Side(style='thin', color=MED_GRAY),
    right=Side(style='thin', color=MED_GRAY),
    top=Side(style='thin', color=MED_GRAY),
    bottom=Side(style='medium', color=NAVY),
)

font_title = Font(name='Calibri', size=18, bold=True, color=WHITE)
font_subtitle = Font(name='Calibri', size=11, bold=False, color=LIGHT_GOLD, italic=True)
font_header = Font(name='Calibri', size=10, bold=True, color=WHITE)
font_data = Font(name='Calibri', size=9, color=DARK_GRAY)
font_data_bold = Font(name='Calibri', size=9, bold=True, color=DARK_GRAY)
font_total = Font(name='Calibri', size=10, bold=True, color=WHITE)
font_footer = Font(name='Calibri', size=8, italic=True, color="888888")

fill_navy = PatternFill(start_color=NAVY, end_color=NAVY, fill_type='solid')
fill_dark_blue = PatternFill(start_color=DARK_BLUE, end_color=DARK_BLUE, fill_type='solid')
fill_mid_blue = PatternFill(start_color=MID_BLUE, end_color=MID_BLUE, fill_type='solid')
fill_teal = PatternFill(start_color=TEAL, end_color=TEAL, fill_type='solid')
fill_gold = PatternFill(start_color=GOLD, end_color=GOLD, fill_type='solid')
fill_light_gray = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type='solid')
fill_very_light = PatternFill(start_color=VERY_LIGHT_GRAY, end_color=VERY_LIGHT_GRAY, fill_type='solid')
fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type='solid')
fill_light_blue = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type='solid')

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
align_right = Alignment(horizontal='right', vertical='center')


# ══════════════════════════════════════════════════════════════════
# HELPER: Title Banner
# ══════════════════════════════════════════════════════════════════

def add_title_banner(ws, title: str, subtitle: str, max_col: int):
    """Add a professional navy/gold title banner to the worksheet."""
    ws.row_dimensions[1].height = 45
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
    cell = ws.cell(row=1, column=1)
    cell.value = f"  {title}"
    cell.font = font_title
    cell.fill = fill_navy
    cell.alignment = Alignment(horizontal='left', vertical='center')
    for c in range(2, max_col + 1):
        ws.cell(row=1, column=c).fill = fill_navy

    ws.row_dimensions[2].height = 25
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
    cell = ws.cell(row=2, column=1)
    cell.value = f"  {subtitle}"
    cell.font = font_subtitle
    cell.fill = fill_dark_blue
    cell.alignment = Alignment(horizontal='left', vertical='center')
    for c in range(2, max_col + 1):
        ws.cell(row=2, column=c).fill = fill_dark_blue

    ws.row_dimensions[3].height = 4
    for c in range(1, max_col + 1):
        ws.cell(row=3, column=c).fill = fill_gold
    ws.row_dimensions[4].height = 8


# ══════════════════════════════════════════════════════════════════
# EXCEL SHEET BUILDERS
# ══════════════════════════════════════════════════════════════════

def build_activities_sheet(wb, activities: list, project_name: str) -> None:
    """Build the Schedule & Cost sheet with all activities."""
    ws = wb.create_sheet(title="Schedule & Cost")
    cols = 13
    add_title_banner(ws, "PROJECT SCHEDULE & COST REPORT",
                     f"{project_name}  |  {len(activities)} Activities  |  Generated {datetime.now().strftime('%d-%b-%Y')}", cols)

    headers = ['Activity\nID', 'Activity Name', 'Duration\n(days)', 'Start\nDate', 'Finish\nDate',
               'Budgeted\nCost (PKR)', 'Actual\nCost (PKR)', 'Variance\n(PKR)', 'Progress\n%',
               'Total\nFloat', 'Critical\nPath', 'Status', 'WBS']
    widths = [12, 45, 10, 14, 14, 16, 16, 14, 10, 10, 10, 14, 20]

    hr = 7
    ws.row_dimensions[hr].height = 35
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=ci)
        cell.value = h
        cell.font = font_header
        cell.fill = fill_mid_blue
        cell.alignment = align_center
        cell.border = thick_bottom
        ws.column_dimensions[get_column_letter(ci)].width = w

    r = 8
    total_bud = 0
    total_act = 0
    for idx, act in enumerate(activities):
        ws.row_dimensions[r].height = 22
        variance = (act.budgeted_cost or 0) - (act.actual_cost or 0)
        act_status = act.status.value if hasattr(act.status, 'value') else str(act.status)
        data = [
            act.activity_id, act.name, act.original_duration,
            act.planned_start, act.planned_finish,
            act.budgeted_cost or 0, act.actual_cost or 0,
            variance, act.percent_complete or 0, act.total_float or 0,
            'YES' if act.is_critical else 'No', act_status, ''
        ]
        total_bud += act.budgeted_cost or 0
        total_act += act.actual_cost or 0

        for ci, val in enumerate(data, 1):
            cell = ws.cell(row=r, column=ci)
            if ci in [4, 5] and isinstance(val, (date, datetime)):
                cell.value = val
                cell.number_format = 'DD-MMM-YY'
            elif ci in [6, 7, 8]:
                cell.value = float(val) if val else 0
                cell.number_format = '#,##0.00'
            elif ci == 9:
                cell.value = float(val) if val else 0
                cell.number_format = '0.0'
            elif ci == 10:
                cell.value = float(val) if val else 0
                cell.number_format = '0.0'
            else:
                cell.value = val

            cell.border = thin_border
            cell.fill = fill_white if idx % 2 == 0 else fill_very_light

            if ci == 1:
                cell.font = Font(name='Consolas', size=9, bold=True, color=MID_BLUE)
                cell.alignment = align_center
            elif ci == 2:
                cell.font = font_data
                cell.alignment = align_left
            elif ci == 8:
                cell.font = font_data_bold
                cell.alignment = align_right
                try:
                    if val and float(val) < 0:
                        cell.font = Font(name='Calibri', size=9, bold=True, color=RED)
                    elif val and float(val) > 0:
                        cell.font = Font(name='Calibri', size=9, bold=True, color=GREEN)
                except (ValueError, TypeError):
                    pass
            elif ci == 11:
                cell.font = font_data_bold
                cell.alignment = align_center
                if val == 'YES':
                    cell.font = Font(name='Calibri', size=9, bold=True, color=RED)
            else:
                cell.font = font_data
                cell.alignment = align_center
        r += 1

    # Grand total row
    ws.row_dimensions[r].height = 28
    for ci in range(1, cols + 1):
        cell = ws.cell(row=r, column=ci)
        cell.font = font_total
        cell.fill = fill_navy
        cell.border = thick_bottom
        cell.alignment = align_center
    ws.cell(row=r, column=2).value = f"GRAND TOTAL ({len(activities)} Activities)"
    ws.cell(row=r, column=2).alignment = align_left
    ws.cell(row=r, column=6).value = total_bud
    ws.cell(row=r, column=6).number_format = '#,##0.00'
    ws.cell(row=r, column=7).value = total_act
    ws.cell(row=r, column=7).number_format = '#,##0.00'
    ws.cell(row=r, column=8).value = total_bud - total_act
    ws.cell(row=r, column=8).number_format = '#,##0.00'

    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
    ws.cell(row=r, column=1).value = f"  Generated by ConstructMind AI  |  Moawia Husnain  |  Civil Engineer  |  UET Taxila  |  +923266915744  |  {datetime.now().strftime('%d-%b-%Y %H:%M')}"
    ws.cell(row=r, column=1).font = font_footer
    ws.freeze_panes = ws.cell(row=8, column=3)
    ws.auto_filter.ref = f"A7:{get_column_letter(cols)}{r-2}"


def build_resources_sheet(wb, resources: list, project_name: str) -> None:
    """Build the Resource Dictionary sheet."""
    ws = wb.create_sheet(title="Resources")
    cols = 8
    add_title_banner(ws, "RESOURCE DICTIONARY",
                     f"{project_name}  |  {len(resources)} Resources  |  Generated {datetime.now().strftime('%d-%b-%Y')}", cols)

    headers = ['Resource ID', 'Resource Name', 'Type', 'Max Units', 'Unit', 'Standard Rate\n(PKR/hr)', 'OT Rate\n(PKR/hr)', 'Active']
    widths = [14, 40, 14, 12, 12, 16, 16, 10]
    hr = 7
    ws.row_dimensions[hr].height = 35
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=ci)
        cell.value = h
        cell.font = font_header
        cell.fill = fill_mid_blue
        cell.alignment = align_center
        cell.border = thick_bottom
        ws.column_dimensions[get_column_letter(ci)].width = w

    r = 8
    for idx, res in enumerate(resources):
        ws.row_dimensions[r].height = 20
        rtype = res.resource_type.value if hasattr(res.resource_type, 'value') else str(res.resource_type)
        data = [res.resource_id, res.name, rtype, res.max_units, res.unit_of_measure,
                res.standard_rate, res.overtime_rate, 'Yes' if res.is_active else 'No']
        for ci, val in enumerate(data, 1):
            cell = ws.cell(row=r, column=ci)
            if ci in [6, 7]:
                cell.value = float(val) if val else 0
                cell.number_format = '#,##0.00'
            elif ci == 4:
                cell.value = float(val) if val else 0
                cell.number_format = '0.0'
            else:
                cell.value = val
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = align_center if ci != 2 else align_left
            if ci == 1:
                cell.font = Font(name='Consolas', size=9, bold=True, color=MID_BLUE)
            cell.fill = fill_white if idx % 2 == 0 else fill_very_light
        r += 1
    ws.freeze_panes = ws.cell(row=8, column=2)


def build_boq_sheet(wb, boq_items: list, project_name: str) -> None:
    """Build the Bill of Quantities sheet."""
    ws = wb.create_sheet(title="Bill of Quantities")
    cols = 10
    add_title_banner(ws, "BILL OF QUANTITIES (BOQ)",
                     f"{project_name}  |  {len(boq_items)} Items  |  Generated {datetime.now().strftime('%d-%b-%Y')}", cols)

    headers = ['Item No', 'Description', 'CSI Code', 'Category', 'Qty', 'Unit', 'Unit Rate\n(PKR)', 'Total Amount\n(PKR)', 'AI Category', 'AI Confidence']
    widths = [12, 50, 12, 25, 12, 10, 16, 18, 25, 12]
    hr = 7
    ws.row_dimensions[hr].height = 35
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=ci)
        cell.value = h
        cell.font = font_header
        cell.fill = fill_mid_blue
        cell.alignment = align_center
        cell.border = thick_bottom
        ws.column_dimensions[get_column_letter(ci)].width = w

    r = 8
    grand_total = 0
    for idx, item in enumerate(boq_items):
        ws.row_dimensions[r].height = 22
        data = [item.item_no, item.description, item.csi_code or '', item.csi_category or '',
                item.quantity, item.unit, item.unit_rate, item.total_amount,
                item.ai_category or '', item.ai_confidence or 0]
        grand_total += item.total_amount or 0
        for ci, val in enumerate(data, 1):
            cell = ws.cell(row=r, column=ci)
            if ci in [5, 7, 8]:
                cell.value = float(val) if val else 0
                cell.number_format = '#,##0.00'
            elif ci == 10:
                cell.value = float(val) if val else 0
                cell.number_format = '0%'
            else:
                cell.value = val
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = align_center if ci not in [2, 4, 9] else align_left
            if ci == 1:
                cell.font = Font(name='Consolas', size=9, bold=True, color=MID_BLUE)
            cell.fill = fill_white if idx % 2 == 0 else fill_very_light
        r += 1

    # Total row
    ws.row_dimensions[r].height = 28
    for ci in range(1, cols + 1):
        cell = ws.cell(row=r, column=ci)
        cell.font = font_total
        cell.fill = fill_navy
        cell.border = thick_bottom
        cell.alignment = align_center
    ws.cell(row=r, column=2).value = f"GRAND TOTAL ({len(boq_items)} Items)"
    ws.cell(row=r, column=2).alignment = align_left
    ws.cell(row=r, column=8).value = grand_total
    ws.cell(row=r, column=8).number_format = '#,##0.00'
    ws.freeze_panes = ws.cell(row=8, column=2)


def build_relationships_sheet(wb, relationships: list, activities: list, project_name: str) -> None:
    """Build the Activity Relationships sheet."""
    ws = wb.create_sheet(title="Relationships")
    cols = 5
    add_title_banner(ws, "ACTIVITY RELATIONSHIPS",
                     f"{project_name}  |  {len(relationships)} Links  |  Generated {datetime.now().strftime('%d-%b-%Y')}", cols)

    act_map = {a.id: a for a in activities}
    headers = ['Predecessor', 'Successor', 'Type', 'Lag (days)', 'Notes']
    widths = [40, 40, 18, 12, 30]
    type_names = {'FS': 'Finish-to-Start', 'SS': 'Start-to-Start', 'FF': 'Finish-to-Finish', 'SF': 'Start-to-Finish'}
    hr = 7
    ws.row_dimensions[hr].height = 35
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=hr, column=ci)
        cell.value = h
        cell.font = font_header
        cell.fill = fill_mid_blue
        cell.alignment = align_center
        cell.border = thick_bottom
        ws.column_dimensions[get_column_letter(ci)].width = w

    r = 8
    for idx, rel in enumerate(relationships):
        pred = act_map.get(rel.predecessor_id)
        succ = act_map.get(rel.successor_id)
        data = [
            f"{pred.activity_id} - {pred.name}" if pred else rel.predecessor_id,
            f"{succ.activity_id} - {succ.name}" if succ else rel.successor_id,
            type_names.get(rel.relationship_type, rel.relationship_type),
            rel.lag_days, ''
        ]
        for ci, val in enumerate(data, 1):
            cell = ws.cell(row=r, column=ci)
            cell.value = val
            cell.font = font_data
            cell.border = thin_border
            cell.alignment = align_left if ci in [1, 2] else align_center
            cell.fill = fill_white if idx % 2 == 0 else fill_very_light
        r += 1
    ws.freeze_panes = ws.cell(row=8, column=1)


# ══════════════════════════════════════════════════════════════════
# XER GENERATION
# ══════════════════════════════════════════════════════════════════

def generate_xer_content(project, activities, relationships, resources, wbs_items) -> str:
    """Generate Primavera P6 XER file content from project data."""
    lines = []
    now = datetime.now()
    hrs = project.hours_per_day or 8

    # Header
    lines.append(f"ERMHDR\t20.12\t{now.strftime('%Y-%m-%d')}\tProject\tmoawia\tConstructMind AI\tUSD")
    lines.append("")

    # Currency
    lines.append("%T\tCURRTYPE")
    lines.append("%F\tcurr_id\tdecimal_digit_cnt\tcurr_symbol\tdecimal_symbol\tdigit_group_symbol\tpos_curr_fmt_type\tneg_curr_fmt_type\tcurr_type\tcurr_short_name\tgroup_digit_cnt\tbase_exch_rate")
    lines.append("%R\t1\t2\tPKR\t.\t,\t#1.1\t(#1.1)\tPakistani Rupees\tPKR\t3\t1")

    # Project
    p_start = project.planned_start.strftime('%Y-%m-%d') if project.planned_start else now.strftime('%Y-%m-%d')
    p_end = project.planned_finish.strftime('%Y-%m-%d') if project.planned_finish else (now + timedelta(days=180)).strftime('%Y-%m-%d')
    lines.append("%T\tPROJECT")
    lines.append("%F\tproj_id\tproj_short_name\tname_sep_char\tplan_start_date\tplan_end_date")
    lines.append(f"%R\t1\t{project.code or 'CM'}\t.\t{p_start}\t{p_end}")

    # Calendar
    lines.append("%T\tCALENDAR")
    lines.append("%F\tclndr_id\tdefault_flag\tclndr_name\tproj_id\tday_hr_cnt\tweek_hr_cnt")
    lines.append(f"%R\t1\tY\tStandard 6-Day\t1\t{hrs}\t{hrs * 6}")

    # WBS
    if wbs_items:
        lines.append("%T\tPROJWBS")
        lines.append("%F\twbs_id\tproj_id\tseq_num\twbs_short_name\twbs_name\tparent_wbs_id")
        for idx, w in enumerate(wbs_items):
            lines.append(f"%R\t{idx+1}\t1\t{idx*10}\t{w.code}\t{w.name}\t")

    # Resources
    if resources:
        lines.append("%T\tRSRC")
        lines.append("%F\trsrc_id\trsrc_name\trsrc_short_name\trsrc_type\tactive_flag\tcurr_id")
        for idx, res in enumerate(resources):
            rt = str(getattr(res, 'resource_type', 'labor'))
            rtype = 'RT_Labor' if 'labor' in rt.lower() else ('RT_Mat' if 'material' in rt.lower() else 'RT_Equip')
            lines.append(f"%R\t{idx+1}\t{res.name}\t{res.resource_id}\t{rtype}\tY\t1")

        lines.append("%T\tRSRCRATE")
        lines.append("%F\trsrc_rate_id\trsrc_id\tmax_qty_per_hr\tcost_per_qty\tstart_date")
        for idx, res in enumerate(resources):
            lines.append(f"%R\t{idx+1}\t{idx+1}\t{res.max_units}\t{res.standard_rate}\t{p_start}")

    # Tasks
    lines.append("%T\tTASK")
    lines.append("%F\ttask_id\tproj_id\ttask_code\ttask_name\tstatus_code\ttarget_drtn_hr_cnt\tremain_drtn_hr_cnt\ttarget_start_date\ttarget_end_date\tphys_complete_pct\ttotal_float_hr_cnt\tfree_float_hr_cnt\ttask_type\tduration_type")
    task_id_map = {}
    for idx, act in enumerate(activities):
        pct = act.percent_complete or 0
        sc = 'TK_Complete' if pct >= 100 else ('TK_Active' if pct > 0 else 'TK_NotStart')
        dur_hrs = (act.original_duration or 0) * hrs
        rem_hrs = (act.remaining_duration or act.original_duration or 0) * hrs
        t_start = act.planned_start.strftime('%Y-%m-%d') if act.planned_start else p_start
        t_end = act.planned_finish.strftime('%Y-%m-%d') if act.planned_finish else p_start
        tf = (act.total_float or 0) * hrs
        ff = (act.free_float or 0) * hrs
        task_id_map[act.id] = idx + 1
        lines.append(f"%R\t{idx+1}\t1\t{act.activity_id}\t{act.name}\t{sc}\t{dur_hrs}\t{rem_hrs}\t{t_start}\t{t_end}\t{pct}\t{tf}\t{ff}\tTT_Task\tDT_FixedDrtn")

    # Relationships
    if relationships:
        lines.append("%T\tTASKPRED")
        lines.append("%F\ttask_pred_id\ttask_id\tpred_task_id\tproj_id\tpred_proj_id\tpred_type\tlag_hr_cnt")
        for idx, rel in enumerate(relationships):
            succ_id = task_id_map.get(rel.successor_id, 0)
            pred_id = task_id_map.get(rel.predecessor_id, 0)
            pred_type = f"PR_{rel.relationship_type}"
            lag_hrs = (rel.lag_days or 0) * hrs
            lines.append(f"%R\t{idx+1}\t{succ_id}\t{pred_id}\t1\t1\t{pred_type}\t{lag_hrs}")

    lines.append("%E")
    return "\r\n".join(lines)


# ══════════════════════════════════════════════════════════════════
# API ENDPOINTS
# ══════════════════════════════════════════════════════════════════

@router.get("/export/excel/{export_type}")
async def export_excel(
    project_id: str,
    export_type: str,
    db: AsyncSession = Depends(get_db)
):
    """
    Generate and download a professional Excel report.
    export_type: 'schedule', 'resources', 'boq', 'relationships', 'master' (all combined)
    """
    if not HAS_OPENPYXL:
        raise HTTPException(status_code=500, detail="openpyxl is not installed on the server.")

    # Fetch project
    proj_res = await db.execute(select(Project).where(Project.id == project_id))
    project = proj_res.scalars().first()
    if not project:
        raise HTTPException(status_code=404, detail=f"Project '{project_id}' not found.")

    # Fetch all related data
    act_res = await db.execute(select(Activity).where(Activity.project_id == project_id).order_by(Activity.sort_order.asc()))
    activities = act_res.scalars().all()

    res_res = await db.execute(select(Resource).where(Resource.project_id == project_id))
    resources = res_res.scalars().all()

    boq_res = await db.execute(select(BOQItem).where(BOQItem.project_id == project_id))
    boq_items = boq_res.scalars().all()

    rel_res = await db.execute(select(Relationship).where(Relationship.project_id == project_id))
    relationships = rel_res.scalars().all()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    pname = project.name or 'Project'

    if export_type in ['schedule', 'master']:
        build_activities_sheet(wb, activities, pname)
    if export_type in ['resources', 'master']:
        build_resources_sheet(wb, resources, pname)
    if export_type in ['boq', 'master']:
        build_boq_sheet(wb, boq_items, pname)
    if export_type in ['relationships', 'master']:
        build_relationships_sheet(wb, relationships, activities, pname)

    if len(wb.sheetnames) == 0:
        build_activities_sheet(wb, activities, pname)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = pname.replace(' ', '_').replace('/', '-')[:30]
    filename = f"ConstructMind_{safe_name}_{export_type}_{datetime.now().strftime('%Y%m%d')}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


@router.get("/export/xer")
async def export_xer(
    project_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Generate and download a Primavera P6 XER file."""
    proj_res = await db.execute(select(Project).where(Project.id == project_id))
    project = proj_res.scalars().first()
    if not project:
        raise HTTPException(status_code=404, detail=f"Project '{project_id}' not found.")

    act_res = await db.execute(select(Activity).where(Activity.project_id == project_id).order_by(Activity.sort_order.asc()))
    activities = act_res.scalars().all()

    rel_res = await db.execute(select(Relationship).where(Relationship.project_id == project_id))
    relationships = rel_res.scalars().all()

    res_res = await db.execute(select(Resource).where(Resource.project_id == project_id))
    resources = res_res.scalars().all()

    wbs_res = await db.execute(select(WBS).where(WBS.project_id == project_id))
    wbs_items = wbs_res.scalars().all()

    xer_content = generate_xer_content(project, activities, relationships, resources, wbs_items)

    buf = io.BytesIO(xer_content.encode('utf-8'))
    buf.seek(0)

    safe_name = (project.code or project.name or 'Project').replace(' ', '_').replace('/', '-')[:30]
    filename = f"ConstructMind_{safe_name}_{datetime.now().strftime('%Y%m%d')}.xer"

    return StreamingResponse(
        buf,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
